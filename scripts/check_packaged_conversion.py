"""Smoke-test the frozen app's private PDF-to-Word worker without opening the UI."""
from pathlib import Path
import json
import subprocess
import sys
import tempfile
import time

from docx import Document
from create_conversion_sample import create_pdf


def main():
    if len(sys.argv) != 2:
        raise SystemExit('usage: check_packaged_conversion.py APP_EXECUTABLE')
    executable = Path(sys.argv[1]).resolve()
    with tempfile.TemporaryDirectory(prefix='paperedit-packaged-') as tmp:
        root = Path(tmp)
        source = root / 'input.pdf'
        staged = root / 'staged.docx'
        create_pdf(source)
        proc = subprocess.Popen([str(executable), '--paperedit-convert-worker', str(root)],
                                stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                                stderr=subprocess.PIPE, text=True)
        events = []
        deadline = time.monotonic() + 90
        sent = False
        try:
            while proc.poll() is None and time.monotonic() < deadline:
                ready = root / 'ready.json'
                if ready.exists() and not sent:
                    events.append(json.loads(ready.read_text(encoding='utf-8')))
                    temporary = root / 'command.tmp'
                    temporary.write_text(json.dumps({'type': 'convert', 'staged_path': str(staged)}), encoding='utf-8')
                    temporary.replace(root / 'command.json')
                    sent = True
                time.sleep(.05)
            proc.wait(timeout=1)
        except subprocess.TimeoutExpired:
            proc.kill()
            proc.wait(timeout=10)
            raise RuntimeError('Packaged conversion timed out')
        result_path = root / 'result.json'
        if result_path.exists():
            events.append(json.loads(result_path.read_text(encoding='utf-8')))
        if proc.returncode != 0 or not staged.exists():
            raise RuntimeError(f'packaged worker failed: {events!r}')
        doc = Document(staged)
        text = '\n'.join(p.text for p in doc.paragraphs)
        if 'Packaged conversion test' not in text or '打包转换测试' not in text:
            raise RuntimeError('packaged DOCX is missing editable text')
        if 'Second page' not in text or len(doc.sections) < 2:
            raise RuntimeError('packaged DOCX is missing page structure')
        if not doc.inline_shapes or not any(
            [[c.text for c in r.cells] for r in t.rows] == [['Item', 'Amount'], ['Local', '100']]
            for t in doc.tables):
            raise RuntimeError('packaged DOCX is missing image or editable table')
        print('Packaged conversion passed:', executable)


if __name__ == '__main__':
    main()
