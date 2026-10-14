import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import fitz
import pytest
from PySide6.QtWidgets import QApplication, QMessageBox
from PySide6.QtTest import QTest
from PySide6.QtCore import Qt, QPointF
from pdf_editor.window import EditorWindow
from pdf_editor.canvas import PageCanvas


@pytest.fixture(scope='module')
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def window(app, tmp_path):
    path = tmp_path / 'ui.pdf'
    d = fitz.open()
    p = d.new_page()
    p.insert_text((60, 100), 'Editable example', fontsize=14)
    p.insert_text((60, 180), '界面中文测试', fontsize=16, fontname='china-s')
    d.new_page()
    d.save(path)
    w = EditorWindow()
    w.load_file(path)
    w.show()
    wait_until(app, lambda: bool(w.canvas.spans))
    yield w
    if w.word_export_job and w.word_export_job.running:
        w.word_export_job.cancel()
        w.word_export_job.process.waitForFinished(3000)
    w.session = None
    w.close()
    app.processEvents()


def wait_until(app, predicate, timeout=10000):
    for _ in range(timeout // 10):
        app.processEvents()
        if predicate():
            return
        QTest.qWait(10)
    raise AssertionError('UI state did not arrive')


def test_select_edit_apply_and_undo(window, app):
    span = next(s for s in window.canvas.spans if s.text == 'Editable example')
    window.set_mode('edit')
    window.select_span(span)
    window.text_input.setPlainText('Updated from interface')
    window.width_input.setValue(280)
    window.height_input.setValue(50)
    QTest.mouseClick(window.apply_button, Qt.MouseButton.LeftButton)
    assert 'Updated from interface' in fitz.open(stream=window.session.snapshot(), filetype='pdf')[0].get_text()
    assert window.session.dirty
    window.undo()
    assert 'Editable example' in fitz.open(stream=window.session.snapshot(), filetype='pdf')[0].get_text()
    window.redo()
    assert 'Updated from interface' in fitz.open(stream=window.session.snapshot(), filetype='pdf')[0].get_text()


def test_add_text_search_and_copy(window, app):
    window.set_mode('add')
    window.add_at(60, 250)
    window.text_input.setPlainText('新增中文内容')
    window.width_input.setValue(280)
    window.height_input.setValue(60)
    QTest.mouseClick(window.apply_button, Qt.MouseButton.LeftButton)
    window.search_input.setText('新增中文')
    window.search_next()
    assert window.current_page == 0
    assert window.canvas.search_rects
    wait_until(app, lambda: any(s.text == '新增中文内容' for s in window.canvas.spans))
    window.select_span(next(s for s in window.canvas.spans if s.text == '新增中文内容'))
    window.copy_selected()
    assert QApplication.clipboard().text() == '新增中文内容'


def test_rotated_page_coordinate_mapping(app):
    canvas = PageCanvas()
    canvas.configure_geometry(595, 842, 90, 1.5)
    # Unrotated PDF (60, 100) maps to display (742, 60), then scale.
    p = canvas.to_pdf(QPointF(1113, 90))
    assert abs(p.x - 60) < 0.001
    assert abs(p.y - 100) < 0.001


def test_background_page_results_are_not_stale(window, app):
    window.go_to_page(1)
    window.go_to_page(0)
    wait_until(app, lambda: bool(window.canvas.spans))
    assert window.current_page == 0
    assert any(s.text == 'Editable example' for s in window.canvas.spans)


def test_default_editor_box_does_not_cover_adjacent_table_cell(window, app):
    from pdf_editor.document import DocumentSession
    d = fitz.open()
    p = d.new_page()
    p.insert_text((60, 100), 'A', fontsize=10)
    p.insert_text((90, 100), 'B', fontsize=10)
    window.session = DocumentSession(d.tobytes())
    window.set_mode('edit')
    window.select_span(next(s for s in window.session.spans(0) if s.text == 'A'))
    assert window.edit_rect()[2] < 90
    window.text_input.setPlainText('C')
    window.apply_edit()
    assert 'C' in fitz.open(stream=window.session.snapshot(), filetype='pdf')[0].get_text()
    assert 'B' in fitz.open(stream=window.session.snapshot(), filetype='pdf')[0].get_text()


def test_typing_directly_on_page_and_committing_changes_pdf(window, app):
    span = next(s for s in window.canvas.spans if s.text == 'Editable example')
    # Clicking a text span while reading should enter direct editing.
    window.select_span(span)
    assert window.mode == 'edit'
    assert window.canvas.inline_editor.isVisible()
    window.canvas.inline_editor.setPlainText('Inline change')
    assert window.text_input.toPlainText() == 'Inline change'
    window.canvas.commit_requested.emit()
    assert 'Inline change' in fitz.open(stream=window.session.snapshot(), filetype='pdf')[0].get_text()


def test_pending_text_is_applied_before_export_or_navigation(window, app):
    window.set_mode('edit')
    window.select_span(next(s for s in window.canvas.spans if s.text == 'Editable example'))
    window.text_input.setPlainText('Pending edit')
    assert window.has_pending_edit()
    assert window.flush_pending_edit()
    assert not window.has_pending_edit()
    assert 'Pending edit' in fitz.open(stream=window.session.snapshot(), filetype='pdf')[0].get_text()


def test_pending_invalid_edit_blocks_navigation(window, app):
    window.set_mode('edit')
    window.select_span(next(s for s in window.canvas.spans if s.text == 'Editable example'))
    window.text_input.setPlainText('Overflow ' * 100)
    window.go_to_page(1)
    assert window.current_page == 0
    assert window.has_pending_edit()


def test_word_export_contains_pending_edit_and_keeps_pdf_dirty(window, app, tmp_path, monkeypatch):
    from docx import Document
    monkeypatch.setattr(QMessageBox, 'information', lambda *args, **kwargs: QMessageBox.StandardButton.Ok)
    window.set_mode('edit')
    window.select_span(next(s for s in window.canvas.spans if s.text == 'Editable example'))
    window.text_input.setPlainText('Word export edit')
    destination = tmp_path / 'export.docx'
    assert window._start_word_export(destination)
    wait_until(app, lambda: destination.exists() and window.word_export_job is None, timeout=60000)
    assert 'Word export edit' in '\n'.join(p.text for p in Document(destination).paragraphs)
    assert window.session.dirty


def test_word_export_rejects_scanned_pdf_before_output_dialog(window, app):
    d = fitz.open()
    p = d.new_page()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20), False)
    pix.clear_with(50)
    p.insert_image((20, 20, 200, 200), pixmap=pix)
    from pdf_editor.document import DocumentSession
    window.session = DocumentSession(d.tobytes())
    assert window.export_word_dialog()
    wait_until(app, lambda: window.word_export_job is None)
    assert window.word_export_job is None
    assert 'OCR' in window.warning_label.text()


def test_word_export_rejects_invalid_edit_before_dialog(window, app, monkeypatch):
    window.select_span(next(s for s in window.canvas.spans if s.text == 'Editable example'))
    window.text_input.setPlainText('Overflow ' * 100)
    assert not window.export_word_dialog()
    assert window.has_pending_edit()
    assert window.word_export_job is None


def test_password_protected_word_export_can_be_cancelled(window, app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog
    window.session.password_protected = True
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args, **kwargs: (str(tmp_path/'secret.docx'), ''))
    monkeypatch.setattr(QMessageBox, 'question', lambda *args, **kwargs: QMessageBox.StandardButton.No)
    assert not window.export_word_dialog()
    assert window.word_export_job is None


def test_restricted_extraction_never_starts_worker(window, app):
    window.session.can_extract = False
    assert not window.export_word_dialog()
    assert window.word_export_job is None
    assert '权限' in window.warning_label.text()


def test_real_export_dialog_adds_extension_and_prevents_duplicates(window, app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog
    from docx import Document
    destination = tmp_path / '中文导出'
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args, **kwargs: (str(destination), ''))
    window.select_span(next(s for s in window.canvas.spans if s.text == 'Editable example'))
    window.text_input.setPlainText('Dialog edited text')
    assert window.export_word_dialog()
    job = window.word_export_job
    assert not window.export_word_dialog()
    assert job is window.word_export_job
    assert not window.word_action.isEnabled()
    wait_until(app, lambda: window.word_export_job is None, timeout=30000)
    assert 'Dialog edited text' in '\n'.join(p.text for p in Document(str(destination)+'.docx').paragraphs)
    assert window.session.dirty
    assert window.word_action.isEnabled()


def test_close_can_wait_or_cancel_worker(window, app, tmp_path, monkeypatch):
    from PySide6.QtGui import QCloseEvent
    monkeypatch.setattr(QMessageBox, 'question', lambda *args, **kwargs: QMessageBox.StandardButton.No)
    assert window._start_word_export(tmp_path / 'close.docx')
    job = window.word_export_job
    event = QCloseEvent()
    window.closeEvent(event)
    assert not event.isAccepted()
    assert job.running
    monkeypatch.setattr(QMessageBox, 'question', lambda *args, **kwargs: QMessageBox.StandardButton.Yes)
    event = QCloseEvent()
    window.closeEvent(event)
    assert event.isAccepted()
    assert not job.running
    assert not job.job_dir.exists()


def test_compact_chinese_table_text_can_be_edited_in_place(window, app):
    from pdf_editor.document import DocumentSession
    d = fitz.open()
    p = d.new_page()
    p.insert_text((60, 100), '转账存取', fontname='china-s', fontsize=8)
    p.insert_text((110, 100), '20260910', fontsize=8)
    p.insert_text((60, 114), '下一行文字', fontname='china-s', fontsize=8)
    window.session = DocumentSession(d.tobytes())
    window.set_mode('edit')
    window.select_span(next(s for s in window.session.spans(0) if s.text == '转账存取'))
    window.canvas.inline_editor.setPlainText('编辑成功')
    assert window.apply_edit()
    text = fitz.open(stream=window.session.snapshot(), filetype='pdf')[0].get_text()
    assert '编辑成功' in text
    assert '转账存取' not in text
    assert '20260910' in text and '下一行文字' in text


def test_zoom_restores_the_pending_inline_editor(window, app):
    window.select_span(next(s for s in window.canvas.spans if s.text == 'Editable example'))
    window.canvas.inline_editor.setPlainText('Zoom draft')
    window.set_zoom(110)
    wait_until(app, lambda: bool(window.canvas.spans))
    assert window.canvas.inline_editor.isVisible()
    assert window.canvas.inline_editor.toPlainText() == 'Zoom draft'
    assert window.has_pending_edit()


def test_resize_handle_receives_pointer_events(window, app):
    window.select_span(next(s for s in window.canvas.spans if s.text == 'Editable example'))
    assert window.canvas.resize_handle.isVisible()
    before = window.width_input.value()
    handle = window.canvas.resize_handle
    QTest.mousePress(handle, Qt.MouseButton.LeftButton)
    QTest.mouseMove(handle, handle.rect().center() + __import__('PySide6.QtCore', fromlist=['QPoint']).QPoint(30, 0))
    QTest.mouseRelease(handle, Qt.MouseButton.LeftButton)
    assert window.width_input.value() > before


def test_longer_inline_text_expands_into_free_space(window, app):
    window.select_span(next(s for s in window.canvas.spans if s.text == 'Editable example'))
    before = window.width_input.value()
    window.canvas.inline_editor.setPlainText('A longer editable sentence')
    assert window.width_input.value() > before
    assert window.apply_edit()
    assert 'A longer editable sentence' in fitz.open(stream=window.session.snapshot(), filetype='pdf')[0].get_text()


def test_nonembedded_original_font_is_retained_through_inline_edit(window, app, tmp_path):
    from pdf_editor.document import DocumentSession
    from test_font_fidelity import predefined_font_document
    window.session = DocumentSession(predefined_font_document().tobytes())
    original = window.session.spans(0)[0]
    window.select_span(original)
    window.canvas.inline_editor.setPlainText('转账存取')
    # Commit also exercises the property panel's 0.1 pt position rounding.
    assert window.flush_pending_edit()
    output = tmp_path/'inline-native.pdf'
    window.session.save_as(output)
    changed = next(s for s in DocumentSession.open(output).spans(0) if s.text == '转账存取')
    assert changed.font == original.font
    assert changed.origin == pytest.approx(original.origin, abs=.001)
