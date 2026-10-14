"""Exercise editable Word content and transactional output using synthetic PDFs."""
from importlib import import_module
from pathlib import Path

import fitz
import pytest

from pdf_editor.document import DocumentSession


def conversion():
    return import_module('pdf_editor.conversion')


@pytest.fixture
def word_source():
    with fitz.open() as doc:
        p = doc.new_page(width=595, height=842)
        p.insert_text((60, 80), 'PaperEdit example', fontsize=18)
        p.insert_text((60, 120), '本地转换中文内容', fontsize=14, fontname='china-s')
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 50, 40), False)
        pix.clear_with(120)
        p.insert_image((60, 150, 140, 210), pixmap=pix)
        for x in (60, 210, 360):
            p.draw_line((x, 240), (x, 360))
        for y in (240, 300, 360):
            p.draw_line((60, y), (360, y))
        for x, y, text in [(75, 275, 'Item'), (225, 275, 'Amount'),
                           (75, 335, 'Local'), (225, 335, '100')]:
            p.insert_text((x, y), text, fontsize=12)
        doc.new_page(width=595, height=842).insert_text((60, 100), 'Second page text', fontsize=14)
        return doc.tobytes()


def docx_text(doc):
    return '\n'.join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for r in t.rows for c in r.cells])


def test_editable_chinese_english_images_and_tables(word_source, tmp_path):
    conversion()
    from docx import Document
    out = tmp_path / '可编辑.docx'
    conversion().write_docx(word_source, out)
    doc = Document(out)
    text = docx_text(doc)
    assert 'PaperEdit example' in text
    assert '本地转换中文内容' in text
    assert len(doc.inline_shapes) >= 1
    assert any([[c.text for c in r.cells] for r in t.rows] ==
               [['Item', 'Amount'], ['Local', '100']] for t in doc.tables)


def test_multipage_content_and_page_sections(word_source, tmp_path):
    conversion()
    from docx import Document
    out = tmp_path / 'pages.docx'
    conversion().write_docx(word_source, out)
    doc = Document(out)
    text = docx_text(doc)
    assert 'Second page text' in text
    assert text.index('PaperEdit example') < text.index('Second page text')
    assert len(doc.sections) >= 2


def test_export_keeps_source_history_and_dirty_state(word_source, tmp_path):
    conversion()
    from docx import Document
    source = tmp_path / 'source.pdf'
    source.write_bytes(word_source)
    session = DocumentSession.open(source)
    session.add_text(1, 'Current edited text', 12, (0, 0, 0), (60, 140, 350, 190))
    snapshot = session.snapshot()
    conversion().write_docx(snapshot, tmp_path / 'edited.docx')
    assert 'Current edited text' in docx_text(Document(tmp_path / 'edited.docx'))
    assert source.read_bytes() == word_source
    assert session.snapshot() == snapshot
    assert session.dirty and session.can_undo
    session.undo()
    assert not session.dirty
    session.redo()
    assert session.snapshot() == snapshot


def test_scanned_document_is_rejected(tmp_path):
    c = conversion()
    with fitz.open() as doc:
        p = doc.new_page()
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 30, 30), False)
        pix.clear_with(100)
        p.insert_image((20, 20, 300, 300), pixmap=pix)
        with pytest.raises(c.ConversionError, match='OCR'):
            c.write_docx(doc.tobytes(), tmp_path / 'scanned.docx')
    assert not (tmp_path / 'scanned.docx').exists()


def test_mixed_document_reports_image_page_not_blank_page(word_source):
    with fitz.open(stream=word_source, filetype='pdf') as doc:
        p = doc.new_page()
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20), False)
        pix.clear_with(100)
        p.insert_image((20, 20, 300, 300), pixmap=pix)
        doc.new_page()
        info = conversion().inspect_pdf(doc.tobytes())
    assert info.page_count == 4
    assert info.text_pages == 2
    assert info.image_only_pages == (3,)


def test_bad_pdf_is_rejected():
    c = conversion()
    with pytest.raises(c.ConversionError):
        c.inspect_pdf(b'not a pdf')


@pytest.mark.parametrize('permissions,password,allowed', [
    (fitz.PDF_PERM_PRINT, 'reader', False),
    (fitz.PDF_PERM_PRINT, 'owner', True),
    (fitz.PDF_PERM_COPY, 'reader', True),
])
def test_extraction_permissions_survive_decrypted_snapshot(word_source, tmp_path, permissions, password, allowed):
    path = tmp_path / 'protected.pdf'
    with fitz.open(stream=word_source, filetype='pdf') as doc:
        doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw='reader', owner_pw='owner', permissions=permissions)
    session = DocumentSession.open(path, password)
    assert session.can_extract is allowed
    assert session.password_protected
    assert not fitz.open(stream=session.snapshot(), filetype='pdf').needs_pass


def test_owner_only_protection_disallows_extraction_without_owner_password(word_source, tmp_path):
    path = tmp_path / 'owner-only.pdf'
    with fitz.open(stream=word_source, filetype='pdf') as doc:
        doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw='', owner_pw='owner', permissions=fitz.PDF_PERM_PRINT)
    assert not DocumentSession.open(path).can_extract
    assert DocumentSession.open(path, 'owner').can_extract


def test_publish_to_chinese_path_and_replace_existing_word(word_source, tmp_path):
    conversion()
    from docx import Document
    c = conversion()
    staged = tmp_path / '.staged.docx'
    dest = tmp_path / '输出 文档.docx'
    dest.write_bytes(b'old document')
    c.write_docx(word_source, staged)
    c.publish_docx(staged, dest)
    assert 'PaperEdit example' in docx_text(Document(dest))
    assert not staged.exists()


def test_invalid_output_never_replaces_existing_file(tmp_path):
    c = conversion()
    staged = tmp_path / '.bad.docx'
    staged.write_bytes(b'corrupt zip')
    dest = tmp_path / 'existing.docx'
    dest.write_bytes(b'original')
    with pytest.raises(c.ConversionError):
        c.publish_docx(staged, dest)
    assert dest.read_bytes() == b'original'


def test_non_docx_destination_and_missing_directory_are_rejected(word_source, tmp_path):
    c = conversion()
    stage = tmp_path / 'stage.docx'
    c.write_docx(word_source, stage)
    source = tmp_path / 'source.pdf'
    source.write_bytes(word_source)
    with pytest.raises(c.ConversionError):
        c.publish_docx(stage, source)
    assert source.read_bytes() == word_source
    with pytest.raises(c.ConversionError):
        c.write_docx(word_source, tmp_path / 'missing' / 'out.docx')


def test_page_failure_is_not_reported_as_success(word_source, tmp_path, monkeypatch):
    conversion()
    from pdf2docx import Converter
    c = conversion()
    def broken_page(self, *args, **kwargs):
        if not kwargs.get('ignore_page_error'):
            raise RuntimeError('failed page')
    monkeypatch.setattr(Converter, 'convert', broken_page)
    out = tmp_path / '.output.docx'
    with pytest.raises(c.ConversionError):
        c.write_docx(word_source, out)
    assert not out.exists()
